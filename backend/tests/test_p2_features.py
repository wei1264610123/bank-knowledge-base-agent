"""
P2 功能测试：报表导出 / 智能推荐问题 / 会话自动命名 / 重新生成 / Excel-PPT 解析
对外 LLM 均通过 Mock 规避，不触网、不消耗 API 额度。
"""
import os
import tempfile
from unittest.mock import AsyncMock, MagicMock, patch

from sqlalchemy import select

from app.database import async_session_factory
from app.models.chat import ChatMessage, ChatSession
from app.models.audit import AuditLog
from app.models.feedback import ChatFeedback
from app.services.knowledge_service import KnowledgeService

# ---------- 会话自动命名（P2） ----------


async def test_auto_title_first_message(client, user_token_factory):
    """第一条提问自动生成会话标题（前20字）"""
    token = await user_token_factory()
    headers = {"Authorization": f"Bearer {token}"}
    s = await client.post("/api/chat/sessions", json={}, headers=headers)
    session_id = s.json()["id"]
    assert s.json()["title"] == "新对话"

    fake = MagicMock()
    fake.get_response = AsyncMock(return_value={
        "content": "好的", "references": [], "message_id": "x1"
    })
    long_q = "转账到其他银行需要多少手续费和限额？请详细说明"
    with patch("app.api.chat.ChatService", return_value=fake):
        resp = await client.post(
            "/api/chat/completions",
            json={"session_id": session_id, "message": long_q, "stream": False},
            headers=headers,
        )
    assert resp.status_code == 200, resp.text
    assert resp.json()["content"] == "好的"

    # 标题应变为提问前20字
    lst = await client.get("/api/chat/sessions", headers=headers)
    titles = {x["title"] for x in lst.json()}
    assert long_q[:20] in titles


async def test_auto_title_keeps_renamed(client, user_token_factory):
    """用户手动重命名后，发送消息不覆盖标题"""
    token = await user_token_factory()
    headers = {"Authorization": f"Bearer {token}"}
    s = await client.post("/api/chat/sessions", json={"title": "我的专属会话"}, headers=headers)
    session_id = s.json()["id"]

    fake = MagicMock()
    fake.get_response = AsyncMock(return_value={"content": "好", "references": [], "message_id": "x2"})
    with patch("app.api.chat.ChatService", return_value=fake):
        resp = await client.post(
            "/api/chat/completions",
            json={"session_id": session_id, "message": "信用卡怎么申请？", "stream": False},
            headers=headers,
        )
    assert resp.status_code == 200, resp.text

    lst = await client.get("/api/chat/sessions", headers=headers)
    titles = {x["title"] for x in lst.json()}
    assert "我的专属会话" in titles


# ---------- 智能推荐问题（P2） ----------


async def test_suggested_questions(client, user_token_factory):
    """推荐问题包含热门提问与未解答问题，且 PII 已脱敏"""
    token = await user_token_factory()
    headers = {"Authorization": f"Bearer {token}"}
    me = await client.get("/api/auth/me", headers=headers)
    user_id = me.json()["id"]

    # 清空相关表，避免历史测试数据污染 TOP-N 统计（测试库专用，安全）
    # 注意外键引用顺序：feedback / quality_reviews → messages → sessions
    from sqlalchemy import delete
    async with async_session_factory() as db:
        from app.models.feedback import QuestionRequest
        from app.models.review import QualityReview
        await db.execute(delete(ChatFeedback))
        await db.execute(delete(QualityReview))
        await db.execute(delete(ChatMessage))
        await db.execute(delete(QuestionRequest))
        await db.execute(delete(ChatSession))
        await db.execute(delete(AuditLog))
        await db.commit()

        session = ChatSession(user_id=user_id, title="测试会话")
        db.add(session)
        await db.flush()
        db.add(ChatMessage(session_id=session.id, role="user", content="请问转账限额是多少？"))
        db.add(ChatMessage(session_id=session.id, role="user", content="请问转账限额是多少？"))
        db.add(ChatMessage(session_id=session.id, role="assistant", content="回复", references=[]))
        await db.commit()
        # 未解答共性问题（含手机号，生成后应被脱敏）
        db.add(QuestionRequest(user_id=user_id, content="请联系用户 13812345678 处理"))
        db.add(QuestionRequest(user_id=user_id, content="请联系用户 13812345678 处理"))
        await db.commit()

    resp = await client.get("/api/chat/suggested-questions", headers=headers)
    assert resp.status_code == 200, resp.text
    items = resp.json()
    assert isinstance(items, list)
    assert any("转账限额" in q for q in items)
    # 未解答问题（含手机号）被收录且已脱敏
    masked_rec = next((q for q in items if "联系" in q or "138" in q), None)
    assert masked_rec is not None
    assert "13812345678" not in masked_rec
    assert "138****5678" in masked_rec
    assert len(items) <= 6


# ---------- 重新生成（P2） ----------


async def _get_user_id(client, token):
    me = await client.get("/api/auth/me", headers={"Authorization": f"Bearer {token}"})
    return me.json()["id"]


async def test_regenerate_clears_pair(client, user_token_factory):
    """重新生成：删除末尾 用户问题+AI回答 对"""
    token = await user_token_factory()
    headers = {"Authorization": f"Bearer {token}"}
    user_id = await _get_user_id(client, token)

    async with async_session_factory() as db:
        session = ChatSession(user_id=user_id, title="生成测试")
        db.add(session)
        await db.flush()
        db.add(ChatMessage(session_id=session.id, role="user", content="旧问题"))
        db.add(ChatMessage(session_id=session.id, role="assistant", content="旧答案", references=[]))
        await db.commit()
        session_id = session.id

    resp = await client.delete(f"/api/chat/sessions/{session_id}/regenerate", headers=headers)
    assert resp.status_code == 200, resp.text

    # 重新生成后消息只剩开头保留的（这里应全部清空）
    msgs = await client.get(f"/api/chat/sessions/{session_id}/messages", headers=headers)
    assert msgs.json() == []


async def test_regenerate_blocked_after_feedback(client, user_token_factory):
    """AI 回答已被评价时，禁止重新生成（409）"""
    token = await user_token_factory()
    headers = {"Authorization": f"Bearer {token}"}
    user_id = await _get_user_id(client, token)

    async with async_session_factory() as db:
        session = ChatSession(user_id=user_id, title="生成测试")
        db.add(session)
        await db.flush()
        db.add(ChatMessage(session_id=session.id, role="user", content="问题"))
        db.add(ChatMessage(session_id=session.id, role="assistant", content="答案", references=[]))
        await db.flush()
        from sqlalchemy import select
        ai = (await db.execute(
            select(ChatMessage).where(ChatMessage.session_id == session.id,
                                      ChatMessage.role == "assistant")
        )).scalar_one()
        db.add(ChatFeedback(message_id=ai.id, user_id=user_id, rating="up"))
        await db.commit()
        session_id = session.id

    resp = await client.delete(f"/api/chat/sessions/{session_id}/regenerate", headers=headers)
    assert resp.status_code == 409, resp.text


# ---------- 报表导出（P2） ----------

EXPORT_HEADERS = {
    "questions": ["用户名", "会话", "问题内容", "提问时间"],
    "audit": ["时间", "用户", "操作", "详情", "对象类型", "IP"],
    "feedback": ["时间", "用户", "回答内容", "评价", "原因", "意见"],
}


async def test_export_questions_csv(client, user_token_factory, admin_token):
    """导出问答明细 CSV：带 BOM、含表头与数据、中文不乱码"""
    token = await user_token_factory()
    user_id = await _get_user_id(client, token)

    async with async_session_factory() as db:
        session = ChatSession(user_id=user_id, title="导出会话")
        db.add(session)
        await db.flush()
        db.add(ChatMessage(session_id=session.id, role="user", content="大额转账有什么注意事项？"))
        await db.commit()

    resp = await client.get(
        "/api/admin/export", params={"type": "questions"},
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert resp.status_code == 200, resp.text
    text = resp.text
    assert text.startswith("\ufeff")  # BOM
    assert "用户名" in text and "大额转账有什么注意事项？" in text


async def test_export_audit_and_feedback_csv(client, user_token_factory, admin_token):
    """导出审计日志与反馈记录 CSV"""
    token = await user_token_factory()
    headers = {"Authorization": f"Bearer {token}"}
    user_id = await _get_user_id(client, token)

    async with async_session_factory() as db:
        db.add(AuditLog(username="用户A", action="login", detail="测试", ip="127.0.0.1"))
        await db.commit()

        session = ChatSession(user_id=user_id, title="反馈会话")
        db.add(session)
        await db.flush()
        db.add(ChatMessage(session_id=session.id, role="user", content="这个问题答案如何？"))
        db.add(ChatMessage(session_id=session.id, role="assistant", content="这是我的回答", references=[]))
        await db.flush()
        from sqlalchemy import select
        ai = (await db.execute(
            select(ChatMessage).where(ChatMessage.session_id == session.id,
                                      ChatMessage.role == "assistant")
        )).scalar_one()
        db.add(ChatFeedback(message_id=ai.id, user_id=user_id, rating="up", reason="helpful"))
        await db.commit()

    admin_headers = {"Authorization": f"Bearer {admin_token}"}

    resp_a = await client.get("/api/admin/export", params={"type": "audit"}, headers=admin_headers)
    assert resp_a.status_code == 200, resp_a.text
    assert resp_a.text.startswith("\ufeff")
    assert "login" in resp_a.text

    resp_f = await client.get("/api/admin/export", params={"type": "feedback"}, headers=admin_headers)
    assert resp_f.status_code == 200, resp_f.text
    assert resp_f.text.startswith("\ufeff")
    assert "这是我的回答" in resp_f.text and "up" in resp_f.text


async def test_export_requires_admin(client, user_token_factory):
    """普通用户不能导出报表"""
    token = await user_token_factory()
    resp = await client.get(
        "/api/admin/export", params={"type": "questions"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 403


async def test_export_invalid_type(client, admin_token):
    """非法导出类型返回 400"""
    resp = await client.get(
        "/api/admin/export", params={"type": "hack"},
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert resp.status_code == 400


# ---------- Excel / PPT 解析（P2） ----------


def test_excel_loader():
    """Excel(.xlsx) 加载器：逐工作表生成文本"""
    path = os.path.join(tempfile.gettempdir(), "ut_test_book.xlsx")
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "利率表"
    ws.append(["产品", "年利率"])
    ws.append(["定期一年", "1.75%"])
    wb.save(path)
    try:
        loader = KnowledgeService()._get_loader(path, ".xlsx")
        docs = list(loader.lazy_load())
        assert len(docs) >= 1
        text = docs[0].page_content
        assert "[工作表: 利率表]" in text
        assert "定期一年" in text and "1.75%" in text
    finally:
        os.remove(path)


def test_pptx_loader():
    """PowerPoint(.pptx) 加载器：逐幻灯片生成文本（含表格）"""
    path = os.path.join(tempfile.gettempdir(), "ut_test_slides.pptx")
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "贷款产品介绍"
    body = slide.shapes.placeholders[1]
    body.text_frame.text = "首套房利率 3.8%"
    prs.save(path)
    try:
        loader = KnowledgeService()._get_loader(path, ".pptx")
        docs = list(loader.lazy_load())
        text = "".join(d.page_content for d in docs)
        assert "[幻灯片 1]" in text
        assert "贷款产品介绍" in text
        assert "3.8%" in text
    finally:
        os.remove(path)